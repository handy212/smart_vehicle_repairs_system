from io import BytesIO

import openpyxl
from django.contrib.auth import get_user_model
from django.test import TestCase

from apps.branches.models import Branch
from apps.customers.models import Customer
from apps.data_exchange.importers.customers_vehicles import CustomersVehiclesImporter
from apps.vehicles.models import Vehicle

User = get_user_model()


def _xlsx_from_rows(headers, rows):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    buffer.name = 'test.xlsx'
    return buffer


class CustomersVehiclesImporterTests(TestCase):
    def setUp(self):
        self.staff = User.objects.create_user(
            username='staff@example.com',
            email='staff@example.com',
            password='pass12345',
            role='admin',
        )
        self.branch = Branch.objects.create(
            name='HQ',
            code='HQ',
            is_active=True,
            is_headquarters=True,
            created_by=self.staff,
        )
        self.importer = CustomersVehiclesImporter()

    def test_erp_preview_links_multiple_vehicles_to_one_customer(self):
        buffer = _xlsx_from_rows(
            ['REG NO', 'CUST_NAME', 'MAKE', 'MODEL', 'MFG  BNYEAR', 'ENGVIN', 'TEL_FAXNO'],
            [
                ['GR-1001-A', 'ACME LIMITED', 'TOYOTA', 'COROLLA', 2010, '1HGBH41JXMN109186', '0244111111'],
                ['GR-1002-A', 'ACME LIMITED', 'HONDA', 'CIVIC', 2012, '1HGBH41JXMN109187', '0244111111'],
                ['GR-1003-A', 'ACME LIMITED', 'FORD', 'FOCUS', None, '', '0244111111'],
            ],
        )
        result = self.importer.preview(buffer, {
            'generate_placeholder_vin': True,
            'default_year': 2000,
            'match_existing_customers': True,
        })
        self.assertEqual(result.format_detected, 'erp_vehicles')
        self.assertEqual(result.summary['customers_to_create'], 1)
        self.assertEqual(result.summary['vehicles_to_create'], 3)
        self.assertGreaterEqual(result.summary['missing_vin_placeholders'], 1)

    def test_commit_matches_existing_customer_by_phone(self):
        user = User.objects.create_user(
            username='existing@example.com',
            email='existing@example.com',
            first_name='Jane',
            last_name='Doe',
            phone='0244222222',
            role='customer',
            password='pass12345',
        )
        customer = Customer.objects.create(
            user=user,
            customer_number='CUS-TEST-0001',
            customer_type='individual',
            status='active',
        )

        buffer = _xlsx_from_rows(
            ['REG NO', 'CUST_NAME', 'MAKE', 'MODEL', 'MFG  BNYEAR', 'ENGVIN', 'TEL_FAXNO'],
            [
                ['GT-55-B', 'JANE DOE', 'TOYOTA', 'RAV4', 2018, '1HGBH41JXMN109199', '0244222222'],
            ],
        )
        result = self.importer.commit(buffer, {
            'generate_placeholder_vin': False,
            'match_existing_customers': True,
        })
        self.assertEqual(result.summary['customers_created'], 0)
        self.assertEqual(result.summary['customers_matched'], 1)
        self.assertEqual(result.summary['vehicles_created'], 1)
        vehicle = Vehicle.objects.get(license_plate='GT-55-B')
        self.assertEqual(vehicle.owner_id, customer.id)

    def test_invalid_vin_without_placeholder_is_skipped(self):
        buffer = _xlsx_from_rows(
            ['REG NO', 'CUST_NAME', 'MAKE', 'MODEL', 'MFG  BNYEAR', 'ENGVIN', 'TEL_FAXNO'],
            [
                ['GT-66-C', 'SAM SAMPLE', 'TOYOTA', 'YARIS', 2015, 'BADVIN', '0244333333'],
            ],
        )
        result = self.importer.preview(buffer, {
            'generate_placeholder_vin': False,
            'match_existing_customers': True,
        })
        self.assertEqual(result.summary['vehicles_to_create'], 0)
        self.assertGreaterEqual(result.summary['vehicles_failed'], 1)

    def test_repairs_common_vin_ocr_characters(self):
        # Contains letter O which is illegal in VIN; should repair to 0
        buffer = _xlsx_from_rows(
            ['REG NO', 'CUST_NAME', 'MAKE', 'MODEL', 'MFG  BNYEAR', 'ENGVIN', 'TEL_FAXNO'],
            [
                ['GT-77-D', 'VIN FIX', 'TOYOTA', 'CAMRY', 2016, '1GNDM19WOXB166303', '0244444444'],
            ],
        )
        result = self.importer.preview(buffer, {
            'generate_placeholder_vin': False,
            'match_existing_customers': True,
        })
        self.assertEqual(result.summary['vehicles_to_create'], 1)
        self.assertGreaterEqual(result.summary['vins_repaired'], 1)
        self.assertEqual(result.summary.get('missing_vin_placeholders', 0), 0)

    def test_duplicate_vin_uses_placeholder_for_later_plates(self):
        buffer = _xlsx_from_rows(
            ['REG NO', 'CUST_NAME', 'MAKE', 'MODEL', 'MFG  BNYEAR', 'ENGVIN', 'TEL_FAXNO'],
            [
                ['AA-1-A', 'FLEET CO', 'TOYOTA', 'HILUX', 2015, '1HGBH41JXMN109186', '0200000001'],
                ['AA-2-A', 'FLEET CO', 'TOYOTA', 'HILUX', 2015, '1HGBH41JXMN109186', '0200000001'],
            ],
        )
        result = self.importer.preview(buffer, {
            'generate_placeholder_vin': True,
            'placeholder_vin_on_duplicate': True,
        })
        self.assertEqual(result.summary['customers_to_create'], 1)
        self.assertEqual(result.summary['vehicles_to_create'], 2)
        self.assertEqual(result.summary['duplicate_vin_in_file'], 1)
        self.assertEqual(result.summary['duplicate_vin_placeholders'], 1)
        dup_msgs = [i.message for i in result.issues if i.code == 'duplicate_vin_placeholder']
        self.assertTrue(any('AA-2-A' in msg for msg in dup_msgs))

    def test_duplicate_vin_can_still_skip_when_placeholders_disabled(self):
        buffer = _xlsx_from_rows(
            ['REG NO', 'CUST_NAME', 'MAKE', 'MODEL', 'MFG  BNYEAR', 'ENGVIN', 'TEL_FAXNO'],
            [
                ['AA-1-A', 'FLEET CO', 'TOYOTA', 'HILUX', 2015, '1HGBH41JXMN109186', '0200000001'],
                ['AA-2-A', 'FLEET CO', 'TOYOTA', 'HILUX', 2015, '1HGBH41JXMN109186', '0200000001'],
            ],
        )
        result = self.importer.preview(buffer, {
            'generate_placeholder_vin': True,
            'placeholder_vin_on_duplicate': False,
        })
        self.assertEqual(result.summary['vehicles_to_create'], 1)
        self.assertEqual(result.summary['duplicate_vin_in_file'], 1)
        self.assertTrue(any(i.code == 'duplicate_vin_kept_first' for i in result.issues))

    def test_missing_make_model_without_vin_imports_as_unknown(self):
        buffer = _xlsx_from_rows(
            ['REG NO', 'CUST_NAME', 'MAKE', 'MODEL', 'MFG  BNYEAR', 'ENGVIN', 'TEL_FAXNO', 'APPLICAT'],
            [
                ['ZZ-9-Z', 'NO VIN CO', '', '', 2015, '', '0200000099', 'TRUCK'],
            ],
        )
        result = self.importer.preview(buffer, {
            'generate_placeholder_vin': True,
            'allow_unknown_make_model': True,
            'default_year': 2000,
        })
        self.assertEqual(result.summary['vehicles_to_create'], 1)
        self.assertGreaterEqual(result.summary.get('unknown_make_model', 0), 1)
        self.assertEqual(result.sample_creates[0]['make'], 'Unknown')

    def test_missing_make_model_filled_from_vin_decode(self):
        from unittest.mock import patch

        buffer = _xlsx_from_rows(
            ['REG NO', 'CUST_NAME', 'MAKE', 'MODEL', 'MFG  BNYEAR', 'ENGVIN', 'TEL_FAXNO'],
            [
                ['BB-1-B', 'DECODE ME', '', '', 2018, '1HGBH41JXMN109186', '0200000002'],
            ],
        )
        with patch.object(
            CustomersVehiclesImporter,
            '_decode_vin_cached',
            return_value={'make': 'HONDA', 'model': 'ACCORD', 'year': 2018, 'has_errors': False},
        ):
            result = self.importer.preview(buffer, {
                'decode_vin_for_missing_fields': True,
                'generate_placeholder_vin': False,
            })
        self.assertEqual(result.summary['vehicles_to_create'], 1)
        self.assertEqual(result.summary['vin_decoded_fields'], 1)
        self.assertEqual(result.sample_creates[0]['make'], 'HONDA')
        self.assertEqual(result.sample_creates[0]['model'], 'ACCORD')

    def test_vin_decode_accepts_partial_results_with_vpic_errors(self):
        from unittest.mock import patch

        from apps.vehicles.vin_decoder import VehicleVINDecoder

        buffer = _xlsx_from_rows(
            ['REG NO', 'CUST_NAME', 'MAKE', 'MODEL', 'MFG  BNYEAR', 'ENGVIN', 'TEL_FAXNO'],
            [
                ['BB-2-B', 'GHANA VIN', '', '', 2015, 'AHTKZ3FSX00200475', '0200000008'],
            ],
        )
        with patch.object(
            VehicleVINDecoder,
            'decode_vin',
            return_value=(
                True,
                {
                    'make': 'TOYOTA',
                    'model': 'FORTUNER',
                    'year': 2015,
                    'has_errors': True,
                    'error_message': '1 - Check Digit...',
                    'has_useful_fields': True,
                },
            ),
        ):
            result = self.importer.preview(buffer, {
                'decode_vin_for_missing_fields': True,
                'generate_placeholder_vin': False,
            })
        self.assertEqual(result.summary['vehicles_to_create'], 1)
        self.assertEqual(result.summary['vin_decoded_fields'], 1)
        self.assertEqual(result.sample_creates[0]['make'], 'TOYOTA')
        self.assertEqual(result.sample_creates[0]['model'], 'FORTUNER')

    def test_unknown_sheet_value_triggers_decode_gate(self):
        from unittest.mock import patch

        buffer = _xlsx_from_rows(
            ['REG NO', 'CUST_NAME', 'MAKE', 'MODEL', 'MFG  BNYEAR', 'ENGVIN', 'TEL_FAXNO'],
            [
                ['BB-3-B', 'UNKNOWN ROW', 'Unknown', 'Unknown', 2016, '1HGBH41JXMN109186', '0200000009'],
            ],
        )
        with patch.object(
            CustomersVehiclesImporter,
            '_decode_vin_cached',
            return_value={'make': 'HONDA', 'model': 'CIVIC', 'year': 2016, 'has_useful_fields': True},
        ) as mocked:
            result = self.importer.preview(buffer, {
                'decode_vin_for_missing_fields': True,
                'generate_placeholder_vin': False,
            })
        mocked.assert_called()
        self.assertEqual(result.sample_creates[0]['make'], 'HONDA')
        self.assertEqual(result.sample_creates[0]['model'], 'CIVIC')

    def test_commit_links_multiple_vehicles_to_one_customer(self):
        buffer = _xlsx_from_rows(
            ['REG NO', 'CUST_NAME', 'MAKE', 'MODEL', 'MFG  BNYEAR', 'ENGVIN', 'TEL_FAXNO'],
            [
                ['CC-1-C', 'MULTI OWNER LTD', 'FORD', 'RANGER', 2014, '1FTEX1EPXFKA00001', '0200000003'],
                ['CC-2-C', 'MULTI OWNER LTD', 'FORD', 'ESCAPE', 2016, '1FMCU0G9XGUA00002', '0200000003'],
            ],
        )
        result = self.importer.commit(buffer, {
            'generate_placeholder_vin': False,
            'decode_vin_for_missing_fields': False,
        })
        self.assertEqual(result.summary['customers_created'], 1)
        self.assertEqual(result.summary['vehicles_created'], 2)
        customer = Customer.objects.get(company_name='MULTI OWNER LTD')
        self.assertEqual(customer.vehicles.count(), 2)
        # Business imports must not use the old "Fleet Account" stub
        self.assertNotEqual(customer.user.first_name, 'Fleet')
        self.assertEqual(customer.company_name, 'MULTI OWNER LTD')
        primary = customer.contacts.filter(is_primary=True).first()
        self.assertIsNotNone(primary)
        self.assertNotEqual(primary.first_name, 'Fleet')

    def test_commit_persists_vin_decoded_data_when_decode_fills_fields(self):
        from unittest.mock import patch

        buffer = _xlsx_from_rows(
            ['REG NO', 'CUST_NAME', 'MAKE', 'MODEL', 'MFG  BNYEAR', 'ENGVIN', 'TEL_FAXNO'],
            [
                ['DD-1-D', 'DECODE STORE', '', '', 2018, '1HGBH41JXMN109186', '0200000010'],
            ],
        )
        payload = {
            'make': 'HONDA',
            'model': 'ACCORD',
            'year': 2018,
            'engine_cylinders': 4,
            'fuel_type_primary': 'Gasoline',
            'drive_type': 'FWD',
            'has_useful_fields': True,
        }
        with patch.object(CustomersVehiclesImporter, '_decode_vin_cached', return_value=payload):
            result = self.importer.commit(buffer, {
                'decode_vin_for_missing_fields': True,
                'generate_placeholder_vin': False,
            })
        self.assertEqual(result.summary['vehicles_created'], 1)
        self.assertGreaterEqual(result.summary.get('vin_decoded_stored', 0), 1)
        vehicle = Vehicle.objects.get(license_plate='DD-1-D')
        self.assertEqual(vehicle.make, 'HONDA')
        self.assertEqual(vehicle.vin_decoded_data.get('engine_cylinders'), 4)
        self.assertIsNotNone(vehicle.vin_decoded_at)

    def test_enrich_specs_decodes_even_when_make_model_present(self):
        from unittest.mock import patch

        buffer = _xlsx_from_rows(
            ['REG NO', 'CUST_NAME', 'MAKE', 'MODEL', 'MFG  BNYEAR', 'ENGVIN', 'TEL_FAXNO'],
            [
                ['EE-1-E', 'ENRICH CO', 'TOYOTA', 'COROLLA', 2010, '1HGBH41JXMN109186', '0200000011'],
            ],
        )
        payload = {
            'make': 'HONDA',
            'model': 'ACCORD',
            'year': 2018,
            'engine_cylinders': 4,
            'drive_type': 'FWD',
            'has_useful_fields': True,
        }
        with patch.object(
            CustomersVehiclesImporter, '_decode_vin_cached', return_value=payload
        ) as mocked:
            result = self.importer.commit(buffer, {
                'decode_vin_for_missing_fields': False,
                'decode_vin_enrich_specs': True,
                'generate_placeholder_vin': False,
            })
        mocked.assert_called()
        self.assertEqual(result.summary['vehicles_created'], 1)
        vehicle = Vehicle.objects.get(license_plate='EE-1-E')
        # Spreadsheet make/model preserved; specs JSON stored
        self.assertEqual(vehicle.make, 'TOYOTA')
        self.assertEqual(vehicle.model, 'COROLLA')
        self.assertEqual(vehicle.vin_decoded_data.get('engine_cylinders'), 4)


class DataExchangeRegistryTests(TestCase):
    def test_built_in_modules_registered(self):
        from apps.data_exchange.registry import list_exporters, list_importers

        importer_keys = {item['key'] for item in list_importers()}
        exporter_keys = {item['key'] for item in list_exporters()}
        self.assertIn('customers_vehicles', importer_keys)
        self.assertIn('customers', importer_keys)
        self.assertIn('vehicles', importer_keys)
        self.assertIn('customers_vehicles', exporter_keys)


class NativePreferredTemplateTests(TestCase):
    def test_detects_preferred_native_headers_without_owner(self):
        from apps.data_exchange.utils import detect_format

        headers = [
            'customer_name', 'company_name', 'phone', 'email', 'license_plate',
            'vin', 'make', 'model', 'year', 'exterior_color', 'vehicle_type',
            'engine_type', 'transmission_type',
        ]
        self.assertEqual(detect_format(headers), 'native_customers_vehicles')

    def test_preview_preferred_native_sheet(self):
        importer = CustomersVehiclesImporter()
        buffer = _xlsx_from_rows(
            [
                'customer_name', 'company_name', 'phone', 'email', 'license_plate',
                'vin', 'make', 'model', 'year',
            ],
            [
                [
                    'ACME LIMITED', 'ACME LIMITED', '0244111111', '',
                    'GR-2001-A', '1HGBH41JXMN109186', 'TOYOTA', 'COROLLA', 2010,
                ],
                [
                    'ACME LIMITED', 'ACME LIMITED', '0244111111', '',
                    'GR-2002-A', '1HGBH41JXMN109187', 'HONDA', 'CIVIC', 2012,
                ],
            ],
        )
        result = importer.preview(buffer, {
            'generate_placeholder_vin': True,
            'default_year': 2000,
            'match_existing_customers': True,
        })
        self.assertEqual(result.format_detected, 'native_customers_vehicles')
        self.assertEqual(result.summary['customers_to_create'], 1)
        self.assertEqual(result.summary['vehicles_to_create'], 2)
