"""
Joint customers + vehicles importer.

Supports:
1. ERP migration sheet (REG NO, CUST_NAME, ENGVIN, TEL_FAXNO, ...)
2. Native vehicle template with owner email/id
3. Combined native sheet with customer + vehicle columns
"""
from __future__ import annotations

from collections import defaultdict
from typing import Any

import openpyxl
from django.contrib.auth import get_user_model
from django.db import transaction
from apps.customers.contact_services import (
    apply_business_contact_person_name,
    sync_primary_contact,
)
from apps.customers.models import Customer
from apps.data_exchange.importers.base import (
    BaseImporter,
    ImportCommitResult,
    ImportPreviewResult,
    RowIssue,
)
from apps.data_exchange.utils import (
    clean_value,
    detect_format,
    is_business_name,
    is_valid_vin,
    make_import_email,
    make_placeholder_vin,
    map_engine_type,
    map_transmission_type,
    map_vehicle_type,
    normalize_header,
    normalize_name,
    normalize_phone,
    parse_year,
    repair_vin,
    split_person_name,
)
from apps.vehicles.models import Vehicle
from apps.vehicles.vin_decoder import VehicleVINDecoder, vehicle_model_updates_from_decoded

User = get_user_model()


class CustomersVehiclesImporter(BaseImporter):
    key = 'customers_vehicles'
    label = 'Customers + Vehicles'
    description = (
        'Import customers and vehicles together, linking each vehicle to its owner. '
        'Supports the ERP vehicles database layout and native templates.'
    )
    supports_export = True

    def default_options(self) -> dict[str, Any]:
        return {
            'match_existing_customers': True,
            'update_matched_customers': False,
            'generate_placeholder_vin': True,
            # When the same VIN appears on another plate (in-file or existing),
            # keep the first VIN and assign a plate-based placeholder to the rest.
            'placeholder_vin_on_duplicate': True,
            'default_year': 2000,
            # Import rows with missing make/model as Unknown when VIN decode cannot fill them.
            'allow_unknown_make_model': True,
            'skip_invalid_vin_without_placeholder': True,
            'customer_type_auto': True,
            # When make/model missing and VIN is valid, call NHTSA VPIC.
            # Off by default — each lookup is a remote HTTP call and dominates large imports.
            'decode_vin_for_missing_fields': False,
            # When True, also decode valid VINs that already have make/model so
            # vin_decoded_data (cylinders, drive, fuel, etc.) is stored for the UI.
            # Off by default — slow for large migrations.
            'decode_vin_enrich_specs': False,
            'vin_decode_timeout_seconds': 4.0,
        }

    def preview(self, file_obj, options: dict[str, Any] | None = None) -> ImportPreviewResult:
        opts = {**self.default_options(), **(options or {})}
        plan = self._build_plan(file_obj, opts, persist=False)
        return ImportPreviewResult(
            format_detected=plan['format_detected'],
            summary=plan['summary'],
            issues=plan['issues'],
            sample_creates=plan['sample_creates'],
            options_echo=opts,
        )

    def commit(self, file_obj, options: dict[str, Any] | None = None) -> ImportCommitResult:
        opts = {**self.default_options(), **(options or {})}
        plan = self._build_plan(file_obj, opts, persist=True)
        return ImportCommitResult(
            summary=plan['summary'],
            created_refs=plan['created_refs'],
            issues=plan['issues'],
        )

    def rollback(self, created_refs: dict[str, list[int]]) -> dict[str, Any]:
        deleted = {'customer': 0, 'vehicle': 0, 'user': 0}
        skipped = {'customer': 0, 'vehicle': 0}
        errors: list[str] = []

        vehicle_ids = list(created_refs.get('vehicle') or [])
        customer_ids = list(created_refs.get('customer') or [])

        # Delete vehicles first (FK to customers)
        for vehicle_id in vehicle_ids:
            try:
                vehicle = Vehicle.objects.filter(pk=vehicle_id).first()
                if not vehicle:
                    skipped['vehicle'] += 1
                    continue
                # Safety: do not delete vehicles that gained work-order history after import
                if hasattr(vehicle, 'work_orders') and vehicle.work_orders.exists():
                    skipped['vehicle'] += 1
                    errors.append(f'Vehicle {vehicle_id} has work orders; skipped')
                    continue
                vehicle.delete()
                deleted['vehicle'] += 1
            except Exception as exc:  # noqa: BLE001
                errors.append(f'Vehicle {vehicle_id}: {exc}')

        for customer_id in customer_ids:
            try:
                customer = Customer.objects.filter(pk=customer_id).select_related('user').first()
                if not customer:
                    skipped['customer'] += 1
                    continue
                if customer.vehicles.exists():
                    skipped['customer'] += 1
                    errors.append(f'Customer {customer_id} still has vehicles; skipped')
                    continue
                user = customer.user
                customer.delete()
                deleted['customer'] += 1
                if user and not hasattr(user, 'customer_profile'):
                    # customer delete cascades user? Customer.user is CASCADE from Customer side
                    # User is parent; deleting customer does NOT delete user by default (Customer has FK to User)
                    # Actually: Customer.user is OneToOne CASCADE - on_delete is on Customer side means
                    # when User deleted, Customer deleted. Deleting Customer leaves User.
                    if user.role == 'customer':
                        user.delete()
                        deleted['user'] += 1
            except Exception as exc:  # noqa: BLE001
                errors.append(f'Customer {customer_id}: {exc}')

        return {'deleted': deleted, 'skipped': skipped, 'errors': errors}

    def export_queryset(self, options: dict[str, Any] | None = None):
        return (
            Vehicle.objects.select_related('owner', 'owner__user')
            .order_by('owner_id', 'id')
        )

    # ------------------------------------------------------------------
    # Internal planning / persistence
    # ------------------------------------------------------------------

    def _build_plan(self, file_obj, opts: dict[str, Any], persist: bool) -> dict:
        rows, headers, format_detected = self._load_rows(file_obj)
        issues: list[RowIssue] = []
        sample_creates: list[dict] = []
        created_refs = {'customer': [], 'vehicle': []}

        # Preload existing unique keys
        existing_vins = set(Vehicle.objects.values_list('vin', flat=True))
        existing_plates = {
            clean_value(p).upper() for p in Vehicle.objects.values_list('license_plate', flat=True)
        }
        existing_emails = {
            e.lower() for e in User.objects.filter(role='customer').values_list('email', flat=True)
        }

        # Phone index for matching existing DB customers
        phone_to_customer: dict[str, list[Customer]] = defaultdict(list)
        if opts.get('match_existing_customers'):
            for customer in Customer.objects.select_related('user').all():
                phone = normalize_phone(customer.user.phone or customer.company_phone or '')
                if phone:
                    phone_to_customer[phone].append(customer)

        # Batch-local customer cache + dedupe indexes (one customer → many vehicles)
        customer_cache: dict[str, dict] = {}
        customer_key_by_phone_name: dict[tuple[str, str], str] = {}
        customer_key_by_name: dict[str, str] = {}
        customer_key_by_phone: dict[str, str] = {}
        batch_vins: set[str] = set()
        batch_vin_first_row: dict[str, int] = {}
        batch_plates: set[str] = set()
        vin_decode_cache: dict[str, dict | None] = {}

        summary = {
            'total_rows': 0,
            'blank_rows': 0,
            'customers_to_create': 0,
            'customers_matched': 0,
            'customers_created': 0,
            'vehicles_to_create': 0,
            'vehicles_created': 0,
            'vehicles_skipped': 0,
            'vehicles_failed': 0,
            'duplicate_vin_in_file': 0,
            'duplicate_vin_placeholders': 0,
            'invalid_vin': 0,
            'vins_repaired': 0,
            'missing_vin_placeholders': 0,
            'missing_year_defaults': 0,
            'unknown_make_model': 0,
            'make_model_pending_decode': 0,
            'vin_decoded_fields': 0,
            'vin_decode_failed': 0,
            'vin_decoded_stored': 0,
        }

        parsed_vehicle_rows: list[dict] = []

        for row_number, raw in rows:
            summary['total_rows'] += 1
            mapped = self._map_row(raw, headers, format_detected, row_number, opts)
            if mapped is None:
                summary['blank_rows'] += 1
                continue

            customer_key = self._stable_customer_key(
                mapped,
                customer_cache=customer_cache,
                by_phone_name=customer_key_by_phone_name,
                by_name=customer_key_by_name,
                by_phone=customer_key_by_phone,
                issues=issues,
            )
            mapped['customer_key'] = customer_key

            if customer_key not in customer_cache:
                customer_plan = self._resolve_customer(
                    mapped,
                    opts=opts,
                    existing_emails=existing_emails,
                    phone_to_customer=phone_to_customer,
                    issues=issues,
                )
                customer_plan['phone_norm'] = mapped.get('phone_norm') or ''
                customer_plan['name_norm'] = normalize_name(mapped.get('name') or '')
                customer_cache[customer_key] = customer_plan
                if customer_plan['action'] == 'create':
                    summary['customers_to_create'] += 1
                elif customer_plan['action'] == 'match':
                    summary['customers_matched'] += 1
                elif customer_plan['action'] == 'fail':
                    summary['vehicles_failed'] += 1
                    continue

            customers_only = bool(opts.get('customers_only')) or format_detected == 'native_customers'
            has_vehicle_signal = bool(
                mapped.get('license_plate') or mapped.get('vin') or mapped.get('make')
            )
            if customers_only or not has_vehicle_signal:
                if len(sample_creates) < 25:
                    sample_creates.append({
                        'row_number': row_number,
                        'customer': customer_cache[customer_key].get('display_name'),
                        'customer_action': customer_cache[customer_key]['action'],
                        'email': customer_cache[customer_key].get('email'),
                    })
                continue

            vehicle_plan = self._resolve_vehicle(
                mapped,
                opts=opts,
                existing_vins=existing_vins,
                existing_plates=existing_plates,
                batch_vins=batch_vins,
                batch_vin_first_row=batch_vin_first_row,
                batch_plates=batch_plates,
                vin_decode_cache=vin_decode_cache,
                issues=issues,
                summary=summary,
            )
            if vehicle_plan is None:
                summary['vehicles_skipped'] += 1
                continue

            parsed_vehicle_rows.append({
                'row_number': row_number,
                'customer_key': customer_key,
                'vehicle': vehicle_plan,
                'mapped': mapped,
            })
            summary['vehicles_to_create'] += 1
            if len(sample_creates) < 25:
                sample_creates.append({
                    'row_number': row_number,
                    'customer': customer_cache[customer_key].get('display_name'),
                    'customer_action': customer_cache[customer_key]['action'],
                    'vin': vehicle_plan['vin'],
                    'license_plate': vehicle_plan['license_plate'],
                    'make': vehicle_plan['make'],
                    'model': vehicle_plan['model'],
                    'year': vehicle_plan['year'],
                })

        if not persist:
            return {
                'format_detected': format_detected,
                'summary': summary,
                'issues': issues,
                'sample_creates': sample_creates,
                'created_refs': created_refs,
            }

        # Persist customers then vehicles in small transactions (not one giant atomic).
        # Suppress QBO outbound storms during bulk migration.
        try:
            from apps.quickbooks_online.sync_context import suppress_outbound_qbo_signals
        except Exception:  # noqa: BLE001
            from contextlib import nullcontext as suppress_outbound_qbo_signals

        progress_every = 10
        from apps.data_exchange.services import check_import_cancelled

        with suppress_outbound_qbo_signals():
            for key, plan in customer_cache.items():
                check_import_cancelled(opts)
                if plan['action'] == 'match' and plan.get('customer'):
                    continue
                if plan['action'] != 'create':
                    continue
                try:
                    with transaction.atomic():
                        customer = self._create_customer(plan)
                    plan['customer'] = customer
                    plan['action'] = 'created'
                    created_refs['customer'].append(customer.id)
                    summary['customers_created'] += 1
                    if summary['customers_created'] <= 50:
                        issues.append(RowIssue(
                            row_number=plan.get('row_number') or 0,
                            level='info',
                            entity_type='customer',
                            action='create',
                            code='customer_created',
                            identifier=plan.get('email') or plan.get('display_name') or '',
                            message=f'Created customer #{customer.customer_number}',
                            payload={'customer_id': customer.id},
                        ))
                except Exception as exc:  # noqa: BLE001
                    plan['action'] = 'fail'
                    plan['customer'] = None
                    issues.append(RowIssue(
                        row_number=plan.get('row_number') or 0,
                        level='error',
                        entity_type='customer',
                        action='fail',
                        code='customer_create_failed',
                        identifier=plan.get('display_name') or '',
                        message=str(exc),
                    ))

                if summary['customers_created'] and summary['customers_created'] % progress_every == 0:
                    self._touch_progress(opts, summary, created_refs)

            for item in parsed_vehicle_rows:
                check_import_cancelled(opts)
                # Late VIN decode for missing make/model during commit only
                vehicle_fields = item['vehicle']
                needs_make = self._is_missing_vehicle_field(vehicle_fields.get('make'))
                needs_model = self._is_missing_vehicle_field(vehicle_fields.get('model'))
                decoded = vehicle_fields.get('decoded_payload')
                needs_lookup = (
                    not vehicle_fields.get('vin_is_placeholder')
                    and not decoded
                    and (
                        (
                            opts.get('decode_vin_for_missing_fields', True)
                            and (needs_make or needs_model)
                        )
                        or opts.get('decode_vin_enrich_specs')
                    )
                )
                if needs_lookup:
                    decoded = self._decode_vin_cached(
                        vehicle_fields['vin'], opts=opts, cache=vin_decode_cache
                    )
                if decoded:
                    vehicle_fields['decoded_payload'] = decoded
                    updates = vehicle_model_updates_from_decoded(
                        decoded,
                        current=vehicle_fields,
                        only_blank=True,
                    )
                    if updates:
                        vehicle_fields.update(updates)
                        summary['vin_decoded_fields'] = summary.get('vin_decoded_fields', 0) + 1
                    needs_make = self._is_missing_vehicle_field(vehicle_fields.get('make'))
                    needs_model = self._is_missing_vehicle_field(vehicle_fields.get('model'))

                if needs_make or needs_model:
                    if opts.get('allow_unknown_make_model', True):
                        if needs_make:
                            vehicle_fields['make'] = 'Unknown'
                        if needs_model:
                            vehicle_fields['model'] = 'Unknown'
                        summary['unknown_make_model'] = summary.get('unknown_make_model', 0) + 1
                        issues.append(RowIssue(
                            row_number=item['row_number'],
                            level='warning',
                            entity_type='vehicle',
                            action='create',
                            code='unknown_make_model',
                            identifier=vehicle_fields.get('license_plate') or vehicle_fields.get('vin') or '',
                            message='Make/model still unknown after VIN decode; importing as Unknown',
                        ))
                    else:
                        summary['vehicles_failed'] += 1
                        issues.append(RowIssue(
                            row_number=item['row_number'],
                            level='error',
                            entity_type='vehicle',
                            action='fail',
                            code='missing_make_model',
                            identifier=vehicle_fields.get('license_plate') or vehicle_fields.get('vin') or '',
                            message='Missing make or model after VIN decode attempt',
                        ))
                        continue

                customer = customer_cache[item['customer_key']].get('customer')
                if customer is None:
                    summary['vehicles_failed'] += 1
                    issues.append(RowIssue(
                        row_number=item['row_number'],
                        level='error',
                        entity_type='vehicle',
                        action='fail',
                        code='owner_missing',
                        identifier=vehicle_fields['vin'],
                        message='Owner customer could not be resolved during commit',
                    ))
                    continue
                try:
                    from django.utils import timezone

                    create_kwargs = {
                        'owner': customer,
                        'vin': vehicle_fields['vin'],
                        'make': vehicle_fields['make'],
                        'model': vehicle_fields['model'],
                        'year': vehicle_fields['year'],
                        'license_plate': vehicle_fields['license_plate'],
                        'exterior_color': vehicle_fields.get('exterior_color') or '',
                        'current_mileage': vehicle_fields.get('current_mileage') or 0,
                        'engine_type': vehicle_fields.get('engine_type') or 'gasoline',
                        'transmission_type': vehicle_fields.get('transmission_type') or 'automatic',
                        'vehicle_type': vehicle_fields.get('vehicle_type') or 'other',
                        'trim': vehicle_fields.get('trim') or '',
                        'engine_size': vehicle_fields.get('engine_size') or '',
                        'status': 'active',
                    }
                    decoded_payload = vehicle_fields.get('decoded_payload')
                    if isinstance(decoded_payload, dict):
                        create_kwargs['vin_decoded_data'] = decoded_payload
                        create_kwargs['vin_decoded_at'] = timezone.now()
                        summary['vin_decoded_stored'] = summary.get('vin_decoded_stored', 0) + 1

                    with transaction.atomic():
                        vehicle = Vehicle.objects.create(**create_kwargs)
                    created_refs['vehicle'].append(vehicle.id)
                    summary['vehicles_created'] += 1
                    existing_vins.add(vehicle.vin)
                    existing_plates.add(vehicle.license_plate.upper())
                    if summary['vehicles_created'] <= 50:
                        issues.append(RowIssue(
                            row_number=item['row_number'],
                            level='info',
                            entity_type='vehicle',
                            action='create',
                            code='vehicle_created',
                            identifier=vehicle.vin,
                            message=f'Created vehicle linked to customer #{customer.customer_number}',
                            payload={'vehicle_id': vehicle.id, 'customer_id': customer.id},
                        ))
                except Exception as exc:  # noqa: BLE001
                    summary['vehicles_failed'] += 1
                    issues.append(RowIssue(
                        row_number=item['row_number'],
                        level='error',
                        entity_type='vehicle',
                        action='fail',
                        code='vehicle_create_failed',
                        identifier=vehicle_fields['vin'],
                        message=str(exc),
                    ))

                if summary['vehicles_created'] and summary['vehicles_created'] % progress_every == 0:
                    self._touch_progress(opts, summary, created_refs)

        summary['vehicles_to_create'] = summary['vehicles_created']
        self._touch_progress(opts, summary, created_refs)
        return {
            'format_detected': format_detected,
            'summary': summary,
            'issues': issues,
            'sample_creates': sample_creates,
            'created_refs': created_refs,
        }

    def _touch_progress(self, opts: dict, summary: dict, created_refs: dict) -> None:
        """Optional live progress callback for long commits."""
        from apps.data_exchange.services import CommitCancelled, check_import_cancelled

        check_import_cancelled(opts)
        callback = opts.get('_progress_callback')
        if callable(callback):
            try:
                callback(summary, created_refs)
            except CommitCancelled:
                raise
            except Exception:  # noqa: BLE001
                pass

    def _load_rows(self, file_obj) -> tuple[list[tuple[int, dict]], list[str], str]:
        workbook = openpyxl.load_workbook(file_obj, read_only=True, data_only=True)
        worksheet = workbook.active
        iterator = worksheet.iter_rows(values_only=True)
        raw_headers = next(iterator, None)
        if not raw_headers:
            raise ValueError('Excel file is empty')

        headers = [normalize_header(h) for h in raw_headers]
        format_detected = detect_format(headers)
        if format_detected == 'unknown':
            raise ValueError(
                'Unrecognized spreadsheet layout. Expected ERP columns '
                '(REG NO, CUST_NAME, MAKE, MODEL, ENGVIN, ...) or native vehicle columns '
                '(vin, make, model, year, owner).'
            )

        rows: list[tuple[int, dict]] = []
        for idx, values in enumerate(iterator, start=2):
            row = {}
            for col_index, header in enumerate(headers):
                if not header:
                    continue
                row[header] = values[col_index] if col_index < len(values) else None
            rows.append((idx, row))
        workbook.close()
        return rows, headers, format_detected

    def _map_row(
        self,
        row: dict,
        headers: list[str],
        format_detected: str,
        row_number: int,
        opts: dict,
    ) -> dict | None:
        if not any(clean_value(v) for v in row.values()):
            return None

        if format_detected == 'erp_vehicles':
            name = clean_value(row.get('cust_name'))
            phone = clean_value(row.get('tel_faxno'))
            plate = clean_value(row.get('reg_no')).upper()
            vin = clean_value(row.get('engvin')).upper()
            make = clean_value(row.get('make'))
            model = clean_value(row.get('model'))
            year = parse_year(row.get('mfg_bnyear'), default=None)
            color = clean_value(row.get('colour'))
            engine = map_engine_type(row.get('fueldlv') or '')
            transmission = map_transmission_type(row.get('transtyp') or '')
            vtype = map_vehicle_type(row.get('applicat') or '')
            email = ''
            owner_ref = ''
            company_name = name if is_business_name(name) else ''
        else:
            # Native / combined
            name = clean_value(row.get('customer_name') or row.get('company_name') or '')
            first_name = clean_value(row.get('first_name'))
            last_name = clean_value(row.get('last_name'))
            if first_name or last_name:
                name = f'{first_name} {last_name}'.strip()
            email = clean_value(row.get('email')).lower()
            phone = clean_value(row.get('phone') or row.get('tel_faxno'))
            owner_ref = clean_value(row.get('owner'))
            plate = clean_value(row.get('license_plate') or row.get('reg_no')).upper()
            vin = clean_value(row.get('vin') or row.get('engvin')).upper()
            make = clean_value(row.get('make'))
            model = clean_value(row.get('model'))
            year = parse_year(row.get('year') or row.get('mfg_bnyear'), default=None)
            color = clean_value(row.get('exterior_color') or row.get('colour'))
            engine = map_engine_type(row.get('engine_type') or row.get('fueldlv') or 'gasoline')
            transmission = map_transmission_type(
                row.get('transmission_type') or row.get('transtyp') or 'automatic'
            )
            vtype = map_vehicle_type(row.get('vehicle_type') or row.get('applicat') or '')
            company_name = clean_value(row.get('company_name'))

        if not name and owner_ref:
            name = owner_ref
        if not name and not owner_ref:
            return None

        phone_norm = normalize_phone(phone)
        name_norm = normalize_name(name)
        # Provisional key — finalized by _stable_customer_key for multi-vehicle dedupe
        if email:
            customer_key = f'email:{email}'
        elif owner_ref.isdigit():
            customer_key = f'id:{owner_ref}'
        elif '@' in owner_ref:
            customer_key = f'email:{owner_ref.lower()}'
        elif phone_norm and name_norm:
            customer_key = f'pn:{phone_norm}|{name_norm}'
        elif phone_norm:
            customer_key = f'p:{phone_norm}'
        else:
            customer_key = f'n:{name_norm}'

        return {
            'row_number': row_number,
            'customer_key': customer_key,
            'name': name,
            'name_norm': name_norm,
            'email': email,
            'phone': phone,
            'phone_norm': phone_norm,
            'owner_ref': owner_ref,
            'company_name': company_name,
            'customer_type': clean_value(row.get('customer_type')).lower(),
            'status': clean_value(row.get('status')).lower() or 'active',
            'service_address': clean_value(row.get('service_address')),
            'service_region': clean_value(row.get('service_region')),
            'service_city': clean_value(row.get('service_city')),
            'service_area': clean_value(row.get('service_area')),
            'billing_address': clean_value(row.get('billing_address')),
            'billing_region': clean_value(row.get('billing_region')),
            'billing_city': clean_value(row.get('billing_city')),
            'billing_area': clean_value(row.get('billing_area')),
            'payment_terms': clean_value(row.get('payment_terms')).lower() or 'due_on_receipt',
            'preferred_contact_method': clean_value(row.get('preferred_contact_method')).lower(),
            'license_plate': plate,
            'vin': vin,
            'make': make,
            'model': model,
            'year': year,
            'exterior_color': color,
            'engine_type': engine,
            'transmission_type': transmission,
            'vehicle_type': vtype or 'other',
            'mileage': clean_value(row.get('current_mileage')),
        }

    def _stable_customer_key(
        self,
        mapped: dict,
        *,
        customer_cache: dict,
        by_phone_name: dict[tuple[str, str], str],
        by_name: dict[str, str],
        by_phone: dict[str, str],
        issues: list[RowIssue],
    ) -> str:
        """
        Ensure one customer record is reused across many vehicle rows.

        Priority: email/id → phone+name → same name with matching/empty phone →
        same phone with matching name → provisional key.
        """
        provisional = mapped['customer_key']
        if provisional.startswith('email:') or provisional.startswith('id:'):
            return provisional

        phone_norm = mapped.get('phone_norm') or ''
        name_norm = mapped.get('name_norm') or normalize_name(mapped.get('name') or '')
        row_number = mapped['row_number']

        if phone_norm and name_norm:
            pn = (phone_norm, name_norm)
            if pn in by_phone_name:
                return by_phone_name[pn]

            # Same name already seen: reuse when phone matches or either side lacks phone
            if name_norm in by_name:
                existing_key = by_name[name_norm]
                existing = customer_cache.get(existing_key) or {}
                existing_phone = existing.get('phone_norm') or ''
                if not existing_phone or existing_phone == phone_norm:
                    by_phone_name[pn] = existing_key
                    by_phone.setdefault(phone_norm, existing_key)
                    return existing_key

            # Same phone already seen with the same normalized name
            if phone_norm in by_phone:
                existing_key = by_phone[phone_norm]
                existing = customer_cache.get(existing_key) or {}
                existing_name = existing.get('name_norm') or normalize_name(
                    existing.get('display_name') or ''
                )
                if existing_name == name_norm:
                    by_phone_name[pn] = existing_key
                    by_name.setdefault(name_norm, existing_key)
                    return existing_key
                issues.append(RowIssue(
                    row_number=row_number,
                    level='warning',
                    entity_type='customer',
                    action='create',
                    code='shared_phone_different_name',
                    identifier=phone_norm,
                    message=(
                        f'Phone {phone_norm} also used by "{existing.get("display_name")}"; '
                        f'keeping separate customer for "{mapped.get("name")}"'
                    ),
                ))

            key = f'pn:{phone_norm}|{name_norm}'
            by_phone_name[pn] = key
            by_name.setdefault(name_norm, key)
            by_phone.setdefault(phone_norm, key)
            return key

        if name_norm:
            if name_norm in by_name:
                return by_name[name_norm]
            key = f'n:{name_norm}'
            by_name[name_norm] = key
            if phone_norm:
                by_phone.setdefault(phone_norm, key)
            return key

        if phone_norm:
            if phone_norm in by_phone:
                return by_phone[phone_norm]
            key = f'p:{phone_norm}'
            by_phone[phone_norm] = key
            return key

        return provisional

    def _resolve_customer(
        self,
        mapped: dict,
        opts: dict,
        existing_emails: set[str],
        phone_to_customer: dict[str, list[Customer]],
        issues: list[RowIssue],
    ) -> dict:
        email = mapped['email']
        owner_ref = mapped['owner_ref']
        phone_norm = mapped['phone_norm']
        name = mapped['name']
        row_number = mapped['row_number']

        # Explicit owner email / id (native vehicle template)
        if owner_ref:
            try:
                if '@' in owner_ref:
                    customer = Customer.objects.select_related('user').get(
                        user__email__iexact=owner_ref
                    )
                else:
                    customer = Customer.objects.select_related('user').get(id=int(owner_ref))
                return {
                    'action': 'match',
                    'customer': customer,
                    'display_name': customer.display_name if hasattr(customer, 'display_name') else str(customer),
                    'email': customer.user.email,
                    'row_number': row_number,
                }
            except (Customer.DoesNotExist, ValueError):
                # Fall through to create/match by other keys if name present
                if not name and '@' not in owner_ref:
                    issues.append(RowIssue(
                        row_number=row_number,
                        level='error',
                        entity_type='customer',
                        action='fail',
                        identifier=owner_ref,
                        message=f"Owner '{owner_ref}' not found",
                    ))
                    return {
                        'action': 'fail',
                        'customer': None,
                        'display_name': owner_ref,
                        'email': '',
                        'row_number': row_number,
                    }
                if '@' in owner_ref:
                    email = owner_ref.lower()

        if opts.get('match_existing_customers'):
            if email:
                customer = Customer.objects.select_related('user').filter(user__email__iexact=email).first()
                if customer:
                    issues.append(RowIssue(
                        row_number=row_number,
                        level='info',
                        entity_type='customer',
                        action='match',
                        identifier=email,
                        message='Matched existing customer by email',
                        payload={'customer_id': customer.id},
                    ))
                    return {
                        'action': 'match',
                        'customer': customer,
                        'display_name': str(customer),
                        'email': email,
                        'row_number': row_number,
                    }

            if phone_norm and phone_norm in phone_to_customer:
                candidates = phone_to_customer[phone_norm]
                name_norm = normalize_name(name)
                exact = [
                    c for c in candidates
                    if normalize_name(c.company_name) == name_norm
                    or normalize_name(f'{c.user.first_name} {c.user.last_name}') == name_norm
                ]
                chosen = exact[0] if exact else None
                if chosen:
                    issues.append(RowIssue(
                        row_number=row_number,
                        level='info',
                        entity_type='customer',
                        action='match',
                        identifier=phone_norm,
                        message=f'Matched existing customer by phone (#{chosen.customer_number})',
                        payload={'customer_id': chosen.id},
                    ))
                    return {
                        'action': 'match',
                        'customer': chosen,
                        'display_name': str(chosen),
                        'email': chosen.user.email,
                        'row_number': row_number,
                    }
                issues.append(RowIssue(
                    row_number=row_number,
                    level='warning',
                    entity_type='customer',
                    action='create',
                    code='shared_phone_different_name',
                    identifier=phone_norm,
                    message=(
                        f'Phone {phone_norm} belongs to an existing customer, but the import name '
                        f'"{name}" does not match; creating a separate customer'
                    ),
                ))

        # Create new customer plan
        business = bool(mapped.get('company_name')) or (
            opts.get('customer_type_auto') and is_business_name(name)
        )
        if business:
            # Use the company/org name as the visible identity — never "Fleet Account".
            # (That stub hid the real company on Contacts and looked unlinked.)
            company_name = mapped.get('company_name') or name
            first_name, last_name = split_person_name(company_name)
            customer_type = 'fleet' if 'fleet' in normalize_name(name).lower() else 'business'
        else:
            first_name, last_name = split_person_name(name)
            company_name = ''
            customer_type = 'individual'

        if not email:
            email = make_import_email(name, mapped.get('phone') or '', row_number)
        # Ensure uniqueness against existing + in-batch reserved emails
        base_email = email
        suffix = 1
        while email.lower() in existing_emails:
            local, _, domain = base_email.partition('@')
            email = f'{local}+{suffix}@{domain}'
            suffix += 1
        existing_emails.add(email.lower())

        explicit_type = clean_value(mapped.get('customer_type')).lower()
        if explicit_type in ('individual', 'business', 'fleet'):
            customer_type = explicit_type

        preferred = clean_value(mapped.get('preferred_contact_method')).lower()
        if preferred not in ('email', 'phone', 'sms', 'mail'):
            preferred = 'phone' if mapped.get('phone') else 'email'

        status = clean_value(mapped.get('status')).lower() or 'active'
        if status not in ('active', 'inactive', 'suspended'):
            status = 'active'

        payment_terms = clean_value(mapped.get('payment_terms')).lower() or 'due_on_receipt'

        return {
            'action': 'create',
            'customer': None,
            'display_name': name,
            'email': email,
            'first_name': first_name,
            'last_name': last_name,
            'phone': mapped.get('phone') or '',
            'company_name': company_name,
            'customer_type': customer_type,
            'status': status,
            'service_address': mapped.get('service_address') or '',
            'service_region': mapped.get('service_region') or '',
            'service_city': mapped.get('service_city') or '',
            'service_area': mapped.get('service_area') or '',
            'billing_address': mapped.get('billing_address') or '',
            'billing_region': mapped.get('billing_region') or '',
            'billing_city': mapped.get('billing_city') or '',
            'billing_area': mapped.get('billing_area') or '',
            'payment_terms': payment_terms,
            'preferred_contact_method': preferred,
            'row_number': row_number,
        }

    @staticmethod
    def _is_missing_vehicle_field(value: Any) -> bool:
        text = clean_value(value).upper()
        return text in ('', 'UNKNOWN', 'N/A', 'NA', 'NONE')

    def _decode_vin_cached(
        self,
        vin: str,
        *,
        opts: dict,
        cache: dict[str, dict | None],
    ) -> dict | None:
        if vin in cache:
            return cache[vin]
        try:
            from django.conf import settings

            from apps.data_exchange.services import check_import_cancelled

            check_import_cancelled(opts)
            decoder = VehicleVINDecoder()
            timeout = float(
                opts.get('vin_decode_timeout_seconds')
                or getattr(settings, 'VIN_DECODE_TIMEOUT_SECONDS', 5)
            )
            success, data = decoder.decode_vin(vin, timeout_seconds=timeout)
            if not success or not isinstance(data, dict):
                cache[vin] = None
                return None
            # Accept partial VPIC results even when has_errors is set (common for
            # non-US VINs). Only reject when make and model are both empty.
            if not data.get('make') and not data.get('model'):
                cache[vin] = None
                return None
            cache[vin] = data
            return data
        except Exception:  # noqa: BLE001
            cache[vin] = None
            return None

    def _resolve_vehicle(
        self,
        mapped: dict,
        opts: dict,
        existing_vins: set[str],
        existing_plates: set[str],
        batch_vins: set[str],
        batch_vin_first_row: dict[str, int],
        batch_plates: set[str],
        vin_decode_cache: dict[str, dict | None],
        issues: list[RowIssue],
        summary: dict,
    ) -> dict | None:
        row_number = mapped['row_number']
        plate = mapped['license_plate']
        vin = mapped['vin']
        make = mapped['make']
        model = mapped['model']
        year = mapped['year']

        if not plate:
            issues.append(RowIssue(
                row_number=row_number,
                level='error',
                entity_type='vehicle',
                action='fail',
                code='missing_plate',
                message='Missing license plate / REG NO',
            ))
            summary['vehicles_failed'] += 1
            return None

        plate_key = plate.upper()
        if plate_key in existing_plates or plate_key in batch_plates:
            issues.append(RowIssue(
                row_number=row_number,
                level='error',
                entity_type='vehicle',
                action='skip',
                code='duplicate_plate',
                identifier=plate,
                message=f'License plate {plate} already exists — skipped',
            ))
            return None

        # --- Resolve VIN first (needed for decode + duplicate handling) ---
        vin_is_placeholder = False
        original_vin = vin
        if vin and is_valid_vin(vin):
            pass
        elif vin:
            repaired = repair_vin(vin)
            cleaned = vin.upper().replace(' ', '').replace('-', '')
            if repaired and repaired != cleaned:
                summary['vins_repaired'] += 1
                issues.append(RowIssue(
                    row_number=row_number,
                    level='warning',
                    entity_type='vehicle',
                    action='create',
                    code='vin_repaired',
                    identifier=original_vin,
                    message=f'Repaired VIN "{original_vin}" → "{repaired}" (I/O/Q normalization)',
                ))
                vin = repaired
            elif repaired:
                vin = repaired
            else:
                summary['invalid_vin'] += 1
                if opts.get('generate_placeholder_vin'):
                    issues.append(RowIssue(
                        row_number=row_number,
                        level='warning',
                        entity_type='vehicle',
                        action='create',
                        code='invalid_vin_placeholder',
                        identifier=original_vin,
                        message=f'Invalid VIN "{original_vin}"; generating placeholder from plate',
                    ))
                    vin = make_placeholder_vin(plate, used=batch_vins | existing_vins)
                    vin_is_placeholder = True
                    summary['missing_vin_placeholders'] += 1
                else:
                    issues.append(RowIssue(
                        row_number=row_number,
                        level='error',
                        entity_type='vehicle',
                        action='fail',
                        code='invalid_vin',
                        identifier=original_vin,
                        message=f'Invalid VIN "{original_vin}"',
                    ))
                    summary['vehicles_failed'] += 1
                    return None
        else:
            if opts.get('generate_placeholder_vin'):
                vin = make_placeholder_vin(plate, used=batch_vins | existing_vins)
                vin_is_placeholder = True
                summary['missing_vin_placeholders'] += 1
                issues.append(RowIssue(
                    row_number=row_number,
                    level='warning',
                    entity_type='vehicle',
                    action='create',
                    code='missing_vin_placeholder',
                    identifier=plate,
                    message='Missing VIN; generated placeholder from plate',
                ))
            else:
                issues.append(RowIssue(
                    row_number=row_number,
                    level='error',
                    entity_type='vehicle',
                    action='fail',
                    code='missing_vin',
                    identifier=plate,
                    message='Missing VIN',
                ))
                summary['vehicles_failed'] += 1
                return None

        # Keep first occurrence of a VIN; later duplicates get a placeholder when enabled.
        if vin in existing_vins or vin in batch_vins:
            first_row = batch_vin_first_row.get(vin, 0)
            summary['duplicate_vin_in_file'] += 1
            can_placeholder = (
                opts.get('generate_placeholder_vin')
                and opts.get('placeholder_vin_on_duplicate', True)
            )
            if can_placeholder:
                original_dup_vin = vin
                vin = make_placeholder_vin(plate, used=batch_vins | existing_vins)
                vin_is_placeholder = True
                summary['duplicate_vin_placeholders'] = summary.get('duplicate_vin_placeholders', 0) + 1
                if original_dup_vin in existing_vins:
                    issues.append(RowIssue(
                        row_number=row_number,
                        level='warning',
                        entity_type='vehicle',
                        action='create',
                        code='duplicate_vin_existing_placeholder',
                        identifier=plate,
                        message=(
                            f'VIN {original_dup_vin} already exists; '
                            f'importing plate {plate} with placeholder VIN'
                        ),
                    ))
                else:
                    issues.append(RowIssue(
                        row_number=row_number,
                        level='warning',
                        entity_type='vehicle',
                        action='create',
                        code='duplicate_vin_placeholder',
                        identifier=plate,
                        message=(
                            f'Duplicate VIN {original_dup_vin} (first at row {first_row}); '
                            f'importing plate {plate} with placeholder VIN'
                        ),
                    ))
            elif vin in existing_vins:
                issues.append(RowIssue(
                    row_number=row_number,
                    level='warning',
                    entity_type='vehicle',
                    action='skip',
                    code='duplicate_vin_existing',
                    identifier=vin,
                    message=f'VIN {vin} already exists in the system — skipped (plate {plate})',
                ))
                return None
            else:
                issues.append(RowIssue(
                    row_number=row_number,
                    level='warning',
                    entity_type='vehicle',
                    action='skip',
                    code='duplicate_vin_kept_first',
                    identifier=vin,
                    message=(
                        f'Duplicate VIN {vin}; importing first occurrence '
                        f'(row {first_row}), skipping this row (plate {plate})'
                    ),
                ))
                return None

        # Fill make/model via NHTSA only when those fields are missing.
        # Never decode solely for missing year — default_year covers that.
        missing_make = self._is_missing_vehicle_field(make)
        missing_model = self._is_missing_vehicle_field(model)
        decoded_payload = None
        should_decode = (
            not vin_is_placeholder
            and is_valid_vin(vin)
            and (
                (
                    opts.get('decode_vin_for_missing_fields')
                    and (missing_make or missing_model)
                )
                or opts.get('decode_vin_enrich_specs')
            )
        )
        if should_decode:
            decoded_payload = self._decode_vin_cached(vin, opts=opts, cache=vin_decode_cache)
            if decoded_payload:
                before = {'make': make, 'model': model, 'year': year, 'trim': '', 'engine_size': '',
                          'engine_type': mapped.get('engine_type') or '', 'transmission_type': mapped.get('transmission_type') or ''}
                updates = vehicle_model_updates_from_decoded(
                    decoded_payload, current=before, only_blank=True
                )
                filled = []
                if 'make' in updates:
                    make = updates['make']
                    missing_make = False
                    filled.append('make')
                if 'model' in updates:
                    model = updates['model']
                    missing_model = False
                    filled.append('model')
                if 'year' in updates and year is None:
                    year = updates['year']
                    filled.append('year')
                if filled or opts.get('decode_vin_enrich_specs'):
                    summary['vin_decoded_fields'] += 1
                    issues.append(RowIssue(
                        row_number=row_number,
                        level='info',
                        entity_type='vehicle',
                        action='create',
                        code='vin_decoded_fields',
                        identifier=vin,
                        message=(
                            f'Filled {", ".join(filled)} from VIN decode'
                            if filled else
                            'Stored VIN decode specs for vehicle profile'
                        ),
                        payload={
                            'decoded_fields': filled,
                            'vpic_has_errors': bool(decoded_payload.get('has_errors')),
                        },
                    ))
            else:
                summary['vin_decode_failed'] += 1

        if missing_make or missing_model:
            can_defer_decode = (
                bool(opts.get('decode_vin_for_missing_fields'))
                and not vin_is_placeholder
                and is_valid_vin(vin)
            )
            if can_defer_decode:
                # Keep the row; commit will call NHTSA to fill make/model.
                make = make if not missing_make else 'UNKNOWN'
                model = model if not missing_model else 'UNKNOWN'
                summary['make_model_pending_decode'] = summary.get('make_model_pending_decode', 0) + 1
                issues.append(RowIssue(
                    row_number=row_number,
                    level='warning',
                    entity_type='vehicle',
                    action='create',
                    code='make_model_pending_decode',
                    identifier=plate,
                    message='Missing make/model; will look up from VIN during commit',
                ))
            elif opts.get('allow_unknown_make_model', True):
                # Prefer any partial field; fall back to vehicle_type / Unknown.
                fallback_model = clean_value(mapped.get('vehicle_type') or '')
                if fallback_model in ('', 'other'):
                    fallback_model = 'Unknown'
                else:
                    fallback_model = fallback_model.title()
                if missing_make:
                    make = 'Unknown'
                if missing_model:
                    model = fallback_model if fallback_model != 'Other' else 'Unknown'
                summary['unknown_make_model'] = summary.get('unknown_make_model', 0) + 1
                issues.append(RowIssue(
                    row_number=row_number,
                    level='warning',
                    entity_type='vehicle',
                    action='create',
                    code='unknown_make_model',
                    identifier=plate,
                    message=(
                        'Missing make or model (no valid VIN to look up); '
                        f'importing as {make}/{model}'
                    ),
                ))
            else:
                issues.append(RowIssue(
                    row_number=row_number,
                    level='error',
                    entity_type='vehicle',
                    action='fail',
                    code='missing_make_model',
                    identifier=plate,
                    message=(
                        'Missing make or model'
                        + (
                            ' (VIN decode unavailable or incomplete)'
                            if should_decode else
                            ' (no valid VIN available to look up)'
                        )
                    ),
                ))
                summary['vehicles_failed'] += 1
                return None

        if year is None:
            default_year = opts.get('default_year')
            if default_year:
                year = int(default_year)
                summary['missing_year_defaults'] += 1
                issues.append(RowIssue(
                    row_number=row_number,
                    level='warning',
                    entity_type='vehicle',
                    action='create',
                    code='missing_year_default',
                    identifier=plate,
                    message=f'Missing year; defaulting to {year}',
                ))
            else:
                issues.append(RowIssue(
                    row_number=row_number,
                    level='error',
                    entity_type='vehicle',
                    action='fail',
                    code='missing_year',
                    identifier=plate,
                    message='Missing year',
                ))
                summary['vehicles_failed'] += 1
                return None

        batch_vins.add(vin)
        batch_vin_first_row[vin] = row_number
        batch_plates.add(plate_key)

        mileage = 0
        if mapped.get('mileage'):
            try:
                mileage = int(float(mapped['mileage']))
            except (TypeError, ValueError):
                mileage = 0

        vehicle_plan = {
            'vin': vin,
            'license_plate': plate[:20],
            'make': make[:100],
            'model': model[:100],
            'year': year,
            'exterior_color': (mapped.get('exterior_color') or '')[:50],
            'engine_type': mapped.get('engine_type') or 'gasoline',
            'transmission_type': mapped.get('transmission_type') or 'automatic',
            'vehicle_type': mapped.get('vehicle_type') or 'other',
            'current_mileage': mileage,
            'vin_is_placeholder': vin_is_placeholder,
            'trim': '',
            'engine_size': '',
        }
        if decoded_payload:
            vehicle_plan['decoded_payload'] = decoded_payload
            vehicle_plan.update(vehicle_model_updates_from_decoded(
                decoded_payload, current=vehicle_plan, only_blank=True
            ))
        return vehicle_plan

    def _create_customer(self, plan: dict) -> Customer:
        user = User.objects.create_user(
            username=plan['email'],
            email=plan['email'],
            first_name=plan.get('first_name') or 'Unknown',
            last_name=plan.get('last_name') or 'Customer',
            phone=clean_value(plan.get('phone'))[:20],
            role='customer',
        )
        company_name = plan.get('company_name') or ''
        customer = Customer.objects.create(
            user=user,
            company_name=company_name,
            customer_type=plan.get('customer_type') or 'individual',
            status=plan.get('status') or 'active',
            preferred_contact_method=plan.get('preferred_contact_method') or (
                'phone' if plan.get('phone') else 'email'
            ),
            service_address=plan.get('service_address') or None,
            service_region=plan.get('service_region') or None,
            service_city=plan.get('service_city') or None,
            service_area=plan.get('service_area') or None,
            billing_address=plan.get('billing_address') or '',
            billing_region=plan.get('billing_region') or '',
            billing_city=plan.get('billing_city') or '',
            billing_area=plan.get('billing_area') or '',
            payment_terms=plan.get('payment_terms') or 'due_on_receipt',
            # For business/fleet imports, contact person defaults to the org name
            # until a real person is added later.
            contact_person_name=company_name if company_name else '',
        )
        apply_business_contact_person_name(
            customer,
            first_name=plan.get('first_name'),
            last_name=plan.get('last_name'),
            contact_person_name=company_name or None,
        )
        customer.save(update_fields=['contact_person_name'])
        sync_primary_contact(customer)
        return customer
